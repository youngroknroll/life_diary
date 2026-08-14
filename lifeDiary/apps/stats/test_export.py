"""엑셀 내보내기 — 손으로 쓰던 시트와 같은 모양인가.

만든 파일을 다시 열어 검사한다. 조립 코드가 무엇을 넣었다고 주장하는지가
아니라 파일에 실제로 무엇이 들어갔는지를 본다.

기준은 `월간_소비시간_기록_시트.xlsx` 의 둘째 시트다 — 하루 24×6 격자,
값은 팔레트 코드 1~5, 왼쪽에 카테고리별 일기와 합계.
"""

from datetime import date
from io import BytesIO

import pytest
from openpyxl import load_workbook

from apps.dashboard.models import TimeBlock
from apps.stats.export import build_monthly_workbook
from apps.tags.models import Category, Tag


@pytest.fixture
def user(make_user):
    return make_user(username="exportuser")


def make_tag(user, name, slug):
    return Tag.objects.create(
        user=user, name=name, category=Category.objects.get(slug=slug)
    )


def record(user, tag, on_date, slots=6, first_slot=0):
    for offset in range(slots):
        TimeBlock.objects.create(
            user=user, date=on_date, slot_index=first_slot + offset, tag=tag
        )


def reopen(workbook):
    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return load_workbook(stream)


def find_cell(sheet, text):
    """하루가 9열씩 옆으로 놓이므로 모든 열을 뒤진다."""
    for row in range(1, sheet.max_row + 1):
        for col in range(1, sheet.max_column + 1):
            value = sheet.cell(row=row, column=col).value
            if isinstance(value, str) and text in value:
                return row, col
    return None


def day_block(sheet, label):
    """하루 덩어리의 (머리글 행, 시작 열)."""
    found = find_cell(sheet, label)
    assert found, f"{label} 을 찾지 못했다"
    return found


class TestLayout:
    def test_one_sheet_named_for_the_month(self, user):
        make_tag(user, "집중", "investment")

        book = reopen(build_monthly_workbook(user, 2026, 8))

        assert book.sheetnames == ["2026-08"]

    def test_a_day_block_carries_a_palette_legend(self, user):
        """팔레트 1~5. 원본 시트의 범례와 같은 자리."""
        make_tag(user, "집중", "investment")

        sheet = reopen(build_monthly_workbook(user, 2026, 8)).active
        row, col = day_block(sheet, "8월 1일")

        assert sheet.cell(row=row, column=col + 1).value == "팔레트"
        assert [sheet.cell(row=row, column=col + 3 + i).value for i in range(5)] == [
            1, 2, 3, 4, 5
        ]

    def test_minute_headers_run_from_zero_to_sixty(self, user):
        make_tag(user, "집중", "investment")

        sheet = reopen(build_monthly_workbook(user, 2026, 8)).active
        row, col = find_cell(sheet, "소비시간 별 한 줄 일기")

        assert [sheet.cell(row=row, column=col + 1 + i).value for i in range(7)] == [
            "00", "10", "20", "30", "40", "50", "60"
        ]

    def test_the_grid_is_twenty_four_hours_of_ten_minute_slots(self, user):
        make_tag(user, "집중", "investment")

        sheet = reopen(build_monthly_workbook(user, 2026, 8)).active
        head, col = day_block(sheet, "8월 1일")
        top = head + 2

        hours = [sheet.cell(row=top + h, column=col + 1).value for h in range(24)]
        assert hours[0] == "00"
        assert hours[-1] == "23"


class TestGridValues:
    def test_a_recorded_slot_carries_its_palette_code(self, user):
        """수면이 1번, 투자가 5번. 원본 팔레트 순서다."""
        sleep = make_tag(user, "잠", "sleep")
        record(user, sleep, date(2026, 8, 1), slots=6, first_slot=0)

        sheet = reopen(build_monthly_workbook(user, 2026, 8)).active
        head, col = day_block(sheet, "8월 1일")
        top = head + 2

        assert [sheet.cell(row=top, column=col + 2 + s).value for s in range(6)] == [1] * 6

    def test_investment_is_the_fifth_code(self, user):
        focus = make_tag(user, "집중", "investment")
        record(user, focus, date(2026, 8, 1), slots=6, first_slot=0)

        sheet = reopen(build_monthly_workbook(user, 2026, 8)).active
        head, col = day_block(sheet, "8월 1일")
        top = head + 2

        assert sheet.cell(row=top, column=col + 2).value == 5

    def test_an_empty_slot_stays_empty(self, user):
        focus = make_tag(user, "집중", "investment")
        record(user, focus, date(2026, 8, 1), slots=3, first_slot=0)

        sheet = reopen(build_monthly_workbook(user, 2026, 8)).active
        head, col = day_block(sheet, "8월 1일")
        top = head + 2

        assert sheet.cell(row=top, column=col + 5).value is None


class TestDaySections:
    def test_tag_names_go_in_the_diary_column(self, user):
        """격자에는 코드만 남는다. 태그 이름은 여기서만 살아남는다."""
        focus = make_tag(user, "집중", "investment")
        study = make_tag(user, "학습", "investment")
        record(user, focus, date(2026, 8, 1), slots=6, first_slot=0)
        record(user, study, date(2026, 8, 1), slots=6, first_slot=10)

        sheet = reopen(build_monthly_workbook(user, 2026, 8)).active
        head, col = day_block(sheet, "8월 1일")
        row = head + 2 + 4 * 4        # 투자시간은 다섯 번째 섹션

        assert sheet.cell(row=row + 1, column=col).value == "집중, 학습"

    def test_each_category_reports_its_hours(self, user):
        focus = make_tag(user, "집중", "investment")
        record(user, focus, date(2026, 8, 1), slots=18, first_slot=0)

        sheet = reopen(build_monthly_workbook(user, 2026, 8)).active
        head, col = day_block(sheet, "8월 1일")
        row = head + 2 + 4 * 4

        assert sheet.cell(row=row + 3, column=col).value == 3.0

    def test_self_care_is_not_a_section(self, user):
        """원본에는 있었지만 앱에 없는 카테고리다 (사용자 결정)."""
        make_tag(user, "집중", "investment")

        sheet = reopen(build_monthly_workbook(user, 2026, 8)).active

        assert find_cell(sheet, "Self Care") is None


class TestCalendarWeeks:
    def test_a_week_keeps_the_days_before_the_first_record(self, user):
        """2026-08-05 는 수요일. 그 주 월·화 자리가 남아야 한다."""
        focus = make_tag(user, "집중", "investment")
        record(user, focus, date(2026, 8, 5), slots=6)

        sheet = reopen(build_monthly_workbook(user, 2026, 8)).active

        assert find_cell(sheet, "8월 3일") is not None   # 월
        assert find_cell(sheet, "8월 4일") is not None   # 화
        assert find_cell(sheet, "8월 5일") is not None   # 수

    def test_the_first_week_starts_on_monday(self, user):
        make_tag(user, "집중", "investment")

        sheet = reopen(build_monthly_workbook(user, 2026, 8)).active

        # 2026-08-01 은 토요일. 첫 주 블록은 1일과 2일만 담는다.
        assert find_cell(sheet, "08/01 ~ 08/02") is not None


class TestChartTable:
    def test_a_row_per_day_with_a_column_per_category(self, user):
        focus = make_tag(user, "집중", "investment")
        record(user, focus, date(2026, 8, 1), slots=18)

        sheet = reopen(build_monthly_workbook(user, 2026, 8)).active
        header, _ = find_cell(sheet, "그래프용 데이터")

        assert sheet.cell(row=header, column=2).value == "수면시간"
        assert sheet.cell(row=header, column=6).value == "투자시간"
        assert sheet.cell(row=header + 1, column=1).value == "1일"
        assert sheet.cell(row=header + 1, column=6).value == 3.0
        assert sheet.cell(row=header + 31, column=1).value == "31일"

    def test_the_month_total_row_sums_the_days(self, user):
        make_tag(user, "집중", "investment")

        sheet = reopen(build_monthly_workbook(user, 2026, 8)).active
        header, _ = find_cell(sheet, "그래프용 데이터")
        total = header + 32

        assert sheet.cell(row=total, column=1).value == "달 합계"
        assert str(sheet.cell(row=total, column=2).value).startswith("=SUM(")


class TestCharts:
    def test_one_trend_line_per_category_and_one_doughnut(self, user):
        make_tag(user, "집중", "investment")

        sheet = reopen(build_monthly_workbook(user, 2026, 8)).active
        kinds = sorted(type(chart).__name__ for chart in sheet._charts)

        assert kinds == ["DoughnutChart"] + ["LineChart"] * 5
