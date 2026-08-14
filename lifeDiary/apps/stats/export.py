"""사용자가 수기로 쓰던 `월간_소비시간_기록_시트.xlsx` 둘째 시트의 재현.

2026-08-13 사용자 결정으로 원본과 달라진 곳 —
- `Self Care` 섹션 제외. 앱에 없는 카테고리다
- 주는 달력 주(월~일). 원본은 기록 시작일부터 7일씩 굴렀다
- 한 줄 일기 칸에 태그 이름. 격자에는 코드만 남아 태그가 사라진다
"""

from calendar import monthrange
from datetime import date

from django.utils.translation import gettext, pgettext
from openpyxl import Workbook
from openpyxl.chart import DoughnutChart, LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from apps.core.utils import HOURS_PER_DAY, MINUTES_PER_HOUR, MINUTES_PER_SLOT, SLOTS_PER_HOUR
from apps.dashboard.repositories import TimeBlockRepository
from apps.tags.repositories import CategoryRepository


_time_block_repo = TimeBlockRepository()
_category_repo = CategoryRepository()

DAYS_PER_WEEK = 7
DAY_COLUMNS = 9
DAYS_PER_ROW = 4
GRID_COLUMNS = SLOTS_PER_HOUR
SECTION_ROWS = 4         # 라벨 · 일기 · 여백 · 합계

THIN = Side(style="thin", color="D0D5D2")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center")


def build_monthly_workbook(user, year, month):
    book = Workbook()
    sheet = book.active
    sheet.title = f"{year}-{month:02d}"

    days = _days_of_month(year, month)
    categories = _categories()
    grid, tags_by_day = _read_month(user, days)

    row = 1
    for week in _calendar_weeks(days):
        row = _write_week(sheet, week, grid, tags_by_day, categories, row)

    _write_chart_table(sheet, days, grid, categories, row + 1)
    _size_columns(sheet)
    return book


def _days_of_month(year, month):
    total = monthrange(year, month)[1]
    return [date(year, month, day) for day in range(1, total + 1)]


def _categories():
    """수면부터 투자까지 — 원본 시트의 팔레트 1~5 순서.

    앱의 `display_order` 는 투자가 1번이라 정확히 뒤집는다.
    """
    return list(reversed(list(_category_repo.find_all())))


def _calendar_weeks(days):
    """달력 주(월~일). 달 밖의 날은 `None` 으로 남긴다.

    수요일에 기록을 시작해도 월·화 자리가 비어서 남는다 — 자리를 지우면
    요일이 밀려 "안 했다"가 "없다"로 읽힌다.
    """
    weeks = []
    current = [None] * DAYS_PER_WEEK
    for day in days:
        current[day.weekday()] = day
        if day.weekday() == DAYS_PER_WEEK - 1:
            weeks.append(current)
            current = [None] * DAYS_PER_WEEK
    if any(current):
        weeks.append(current)
    return weeks


def _read_month(user, days):
    grid = {day: [[None] * GRID_COLUMNS for _ in range(HOURS_PER_DAY)] for day in days}
    tags = {day: {} for day in days}

    for block in _time_block_repo.find_by_month(user, days[0], days[-1]):
        if not (block.tag and block.tag.category_id) or block.date not in grid:
            continue
        hour, slot = divmod(block.slot_index, SLOTS_PER_HOUR)
        if hour >= HOURS_PER_DAY:
            continue
        grid[block.date][hour][slot] = block.tag.category_id
        tags[block.date].setdefault(block.tag.category_id, set()).add(block.tag.name)

    return grid, tags


def _write_week(sheet, week, grid, tags_by_day, categories, row):
    sheet.cell(row=row, column=1, value=_week_label(week)).font = Font(bold=True)
    row += 2

    for start in range(0, DAYS_PER_WEEK, DAYS_PER_ROW):
        chunk = week[start : start + DAYS_PER_ROW]
        if not any(chunk):
            continue
        for offset, day in enumerate(chunk):
            _write_day(
                sheet, day, grid, tags_by_day, categories,
                row, 1 + offset * DAY_COLUMNS,
            )
        row += HOURS_PER_DAY + 5
    return row + 1


def _week_label(week):
    present = [day for day in week if day]
    return gettext("%(from)s ~ %(to)s") % {
        "from": present[0].strftime("%m/%d"), "to": present[-1].strftime("%m/%d")
    }


def _write_day(sheet, day, grid, tags_by_day, categories, row, col):
    if day is None:
        return

    sheet.cell(row=row, column=col, value=_day_label(day)).font = Font(bold=True)
    sheet.cell(row=row, column=col + 1, value=gettext("팔레트")).font = Font(size=9)
    for index, category in enumerate(categories):
        cell = sheet.cell(row=row, column=col + 3 + index, value=index + 1)
        cell.fill = PatternFill("solid", fgColor=category.color.lstrip("#"))
        cell.alignment = CENTER

    sheet.cell(row=row + 1, column=col, value=gettext("소비시간 별 한 줄 일기"))
    for index, minute in enumerate(range(0, 70, 10)):
        sheet.cell(row=row + 1, column=col + 1 + index, value=f"{minute:02d}")

    _write_grid(sheet, day, grid, categories, row + 2, col)
    _write_day_sections(sheet, day, grid, tags_by_day, categories, row + 2, col)


def _day_label(day):
    return gettext("%(month)s월 %(day)s일 (%(weekday)s)") % {
        "month": day.month, "day": day.day, "weekday": _weekday(day.weekday())
    }


def _weekday(index):
    """월요일이 0. 리터럴로 적어야 makemessages 가 집어 간다."""
    return (
        pgettext("요일", "월"), pgettext("요일", "화"), pgettext("요일", "수"),
        pgettext("요일", "목"), pgettext("요일", "금"), pgettext("요일", "토"),
        pgettext("요일", "일"),
    )[index]


def _write_grid(sheet, day, grid, categories, row, col):
    code_of = {category.id: index + 1 for index, category in enumerate(categories)}
    color_of = {category.id: category.color.lstrip("#") for category in categories}

    for hour in range(HOURS_PER_DAY):
        label = sheet.cell(row=row + hour, column=col + 1, value=f"{hour:02d}")
        label.font = Font(size=9)
        label.alignment = CENTER
        for slot in range(GRID_COLUMNS):
            cell = sheet.cell(row=row + hour, column=col + 2 + slot)
            cell.border = BORDER
            cell.alignment = CENTER
            category_id = grid[day][hour][slot]
            if category_id in code_of:
                cell.value = code_of[category_id]
                cell.fill = PatternFill("solid", fgColor=color_of[category_id])


def _write_day_sections(sheet, day, grid, tags_by_day, categories, row, col):
    minutes = _minutes_by_category(grid[day])

    for index, category in enumerate(categories):
        top = row + index * SECTION_ROWS
        sheet.cell(row=top, column=col, value=category.display_name).font = Font(
            size=9, bold=True
        )
        names = sorted(tags_by_day[day].get(category.id, ()))
        sheet.cell(row=top + 1, column=col, value=", ".join(names)).font = Font(size=9)
        sheet.cell(
            row=top + 3, column=col,
            value=round(minutes.get(category.id, 0) / MINUTES_PER_HOUR, 2),
        ).font = Font(size=9, bold=True)

    tail = row + len(categories) * SECTION_ROWS
    sheet.cell(row=tail, column=col, value=gettext("기록 합계")).font = Font(size=9)
    sheet.cell(
        row=tail + 1, column=col,
        value=round(sum(minutes.values()) / MINUTES_PER_HOUR, 2),
    ).font = Font(size=9, bold=True)


def _minutes_by_category(day_grid):
    minutes = {}
    for hour_row in day_grid:
        for category_id in hour_row:
            if category_id is not None:
                minutes[category_id] = minutes.get(category_id, 0) + MINUTES_PER_SLOT
    return minutes


def _write_chart_table(sheet, days, grid, categories, row):
    sheet.cell(row=row, column=1, value=gettext("그래프용 데이터")).font = Font(bold=True)
    for index, category in enumerate(categories):
        sheet.cell(row=row, column=2 + index, value=category.display_name)
    sheet.cell(row=row, column=2 + len(categories), value=gettext("합계"))

    for offset, day in enumerate(days):
        minutes = _minutes_by_category(grid[day])
        line = row + 1 + offset
        sheet.cell(row=line, column=1, value=gettext("%(day)s일") % {"day": day.day})
        for index, category in enumerate(categories):
            sheet.cell(
                row=line, column=2 + index,
                value=round(minutes.get(category.id, 0) / MINUTES_PER_HOUR, 2),
            )
        sheet.cell(
            row=line, column=2 + len(categories),
            value=round(sum(minutes.values()) / MINUTES_PER_HOUR, 2),
        )

    total_row = row + 1 + len(days)
    sheet.cell(row=total_row, column=1, value=gettext("달 합계")).font = Font(bold=True)
    for index in range(len(categories)):
        letter = get_column_letter(2 + index)
        sheet.cell(
            row=total_row, column=2 + index,
            value=f"=SUM({letter}{row + 1}:{letter}{total_row - 1})",
        )

    _add_charts(sheet, categories, row, len(days), total_row)


def _add_charts(sheet, categories, header_row, day_count, total_row):
    anchor = total_row + 2
    for index, category in enumerate(categories):
        chart = LineChart()
        chart.title = gettext("%(name)s 추세") % {"name": category.display_name}
        chart.y_axis.title = gettext("시간")
        chart.height, chart.width = 6, 12
        chart.add_data(
            Reference(
                sheet, min_col=2 + index, min_row=header_row,
                max_row=header_row + day_count,
            ),
            titles_from_data=True,
        )
        chart.set_categories(
            Reference(sheet, min_col=1, min_row=header_row + 1,
                      max_row=header_row + day_count)
        )
        sheet.add_chart(chart, f"A{anchor + index * 13}")

    doughnut = DoughnutChart()
    doughnut.title = gettext("이 달 소비시간 비율")
    doughnut.height, doughnut.width = 9, 12
    doughnut.add_data(
        Reference(sheet, min_col=2, max_col=1 + len(categories),
                  min_row=total_row, max_row=total_row),
        from_rows=True,
    )
    doughnut.set_categories(
        Reference(sheet, min_col=2, max_col=1 + len(categories),
                  min_row=header_row, max_row=header_row)
    )
    sheet.add_chart(doughnut, f"J{anchor}")


def _size_columns(sheet):
    for start in range(0, DAYS_PER_ROW * DAY_COLUMNS, DAY_COLUMNS):
        sheet.column_dimensions[get_column_letter(start + 1)].width = 22
        sheet.column_dimensions[get_column_letter(start + 2)].width = 4
        for slot in range(GRID_COLUMNS):
            sheet.column_dimensions[get_column_letter(start + 3 + slot)].width = 3.5
        sheet.column_dimensions[get_column_letter(start + DAY_COLUMNS)].width = 2
